from __future__ import annotations

from io import BytesIO
from typing import Any

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.stats import gaussian_kde

from .models import DiagnosticResult, EventStudyResult, PlaceboResult


def _pyplot():
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "DejaVu Sans"],
            "axes.unicode_minus": False,
            "axes.edgecolor": "#9EABB3",
            "axes.labelcolor": "#263D4B",
            "xtick.color": "#526773",
            "ytick.color": "#526773",
        }
    )
    return plt


def figure_bytes(result: DiagnosticResult, key: str, kind: str) -> bytes:
    plt = _pyplot()
    if key == "placebo_all":
        if len(result.placebos) != 3:
            raise ValueError("All three placebo methods are required for the triptych")
        figure, axes = plt.subplots(1, 3, figsize=(13.2, 4.4), sharey=True, constrained_layout=True)
        for axis, placebo in zip(axes, result.placebos, strict=True):
            draw_placebo(axis, result, placebo, compact=True)
    elif key.startswith("event_"):
        method = key.removeprefix("event_")
        event = next((item for item in result.event_studies if item.method.value == method), None)
        if event is None:
            raise KeyError(key)
        figure, axis = plt.subplots(figsize=(8.2, 5.2), constrained_layout=True)
        draw_event_study(axis, result, event)
    elif key.startswith("placebo_"):
        method = key.removeprefix("placebo_")
        placebo = next((item for item in result.placebos if item.method.value == method), None)
        if placebo is None:
            raise KeyError(key)
        figure, axis = plt.subplots(figsize=(8.2, 5.2), constrained_layout=True)
        draw_placebo(axis, result, placebo)
    else:
        raise KeyError(key)
    stream = BytesIO()
    figure.savefig(stream, format=kind, dpi=300 if kind == "png" else None, bbox_inches="tight")
    plt.close(figure)
    return stream.getvalue()


def draw_event_study(axis: Any, result: DiagnosticResult, event: EventStudyResult) -> None:
    language = result.spec.language.value
    labels = result.spec.labels
    points = [point for point in event.points if point.identifiable and point.estimate is not None]
    x = np.array([point.event_time for point in points])
    y = np.array([point.estimate for point in points], dtype=float)
    low = np.array([point.conf_low for point in points], dtype=float)
    high = np.array([point.conf_high for point in points], dtype=float)
    axis.errorbar(
        x,
        y,
        yerr=np.vstack([y - low, high - y]),
        fmt="o-",
        color="#173B57",
        ecolor="#75A5A3",
        elinewidth=1.2,
        capsize=3,
        markersize=4.8,
        linewidth=1.2,
    )
    axis.axhline(0, color="#596A74", linewidth=0.9)
    axis.axvline(0, color="#A44343", linewidth=1, linestyle="--")
    axis.set_xticks(range(result.spec.window_start, result.spec.window_end + 1))
    axis.set_xlabel("相对政策实施期" if language == "zh" else "Periods relative to treatment")
    y_label = labels.y_axis_zh if language == "zh" else labels.y_axis_en
    unit = labels.unit_zh if language == "zh" else labels.unit_en
    axis.set_ylabel(
        f"{y_label}（{unit}）"
        if language == "zh" and unit
        else f"{y_label} ({unit})"
        if unit
        else y_label
    )
    base_title = labels.title_zh if language == "zh" else labels.title_en
    method_name = "Cohort-saturated" if event.method.value == "saturated" else "Gardner DID2S"
    axis.set_title(f"{base_title} — {method_name}", color="#173B57", pad=12)
    axis.grid(axis="y", color="#DCE3E8", linewidth=0.6, alpha=0.65)
    axis.spines[["top", "right"]].set_visible(False)
    clusters = result.sample.get("clusters", "—")
    note = (
        f"注：基准期为政策前1期；误差线为95%置信区间；标准误按 {clusters} 个聚类计算。"
        if language == "zh"
        else (
            "Notes: The omitted baseline is period -1. Bars show 95% confidence "
            f"intervals; inference uses {clusters} clusters."
        )
    )
    axis.text(
        0, -0.2, note, transform=axis.transAxes, fontsize=8.5, color="#5E707B", va="top", wrap=True
    )


def placebo_title(method: str, language: str) -> str:
    names = {
        "zh": {
            "random_group": "随机处理组安慰剂检验",
            "random_timing": "随机政策时点安慰剂检验",
            "permute_outcome": "结果轨迹置换安慰剂检验",
        },
        "en": {
            "random_group": "Random-group placebo test",
            "random_timing": "Random-timing placebo test",
            "permute_outcome": "Outcome-trajectory placebo test",
        },
    }
    return names[language][method]


def draw_placebo(
    axis: Any, result: DiagnosticResult, placebo: PlaceboResult, compact: bool = False
) -> None:
    language = result.spec.language.value
    values = np.array([draw.estimate for draw in placebo.draws], dtype=float)
    bins = min(35, max(12, int(np.sqrt(len(values)))))
    axis.hist(values, bins=bins, density=True, color="#9FC3BF", edgecolor="#FFFFFF", alpha=0.8)
    if result.spec.show_density and len(np.unique(values)) > 2 and np.std(values) > 0:
        grid = np.linspace(values.min(), values.max(), 240)
        axis.plot(grid, gaussian_kde(values)(grid), color="#173B57", linewidth=1.5)
    axis.axvline(
        placebo.actual_att,
        color="#A44343",
        linewidth=1.5,
        linestyle="--",
        label="真实ATT" if language == "zh" else "Actual ATT",
    )
    axis.axvspan(placebo.quantile_low, placebo.quantile_high, color="#367F82", alpha=0.08)
    axis.set_title(
        placebo_title(placebo.method.value, language),
        color="#173B57",
        fontsize=10.5 if compact else 12.5,
    )
    axis.set_xlabel("安慰剂ATT估计" if language == "zh" else "Placebo ATT estimates")
    if not compact:
        axis.set_ylabel("密度" if language == "zh" else "Density")
    axis.grid(axis="y", color="#DCE3E8", linewidth=0.6, alpha=0.65)
    axis.spines[["top", "right"]].set_visible(False)
    axis.legend(frameon=False, fontsize=8.5)
    annotation = (
        f"经验 p = {placebo.empirical_p_value:.3f}"
        if language == "zh"
        else f"Empirical p = {placebo.empirical_p_value:.3f}"
    )
    axis.text(
        0.98,
        0.96,
        annotation,
        transform=axis.transAxes,
        ha="right",
        va="top",
        fontsize=8.5,
        color="#526773",
    )
    if not compact:
        note = (
            f"注：共获得 {len(values)} 次有效置换；阴影为安慰剂分布的95%区间。"
            if language == "zh"
            else (
                f"Notes: {len(values)} valid permutations. Shading marks the central "
                "95% of the placebo distribution."
            )
        )
        axis.text(
            0,
            -0.2,
            note,
            transform=axis.transAxes,
            fontsize=8.5,
            color="#5E707B",
            va="top",
            wrap=True,
        )


def build_diagnostic_excel(result: DiagnosticResult) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "摘要" if result.spec.language.value == "zh" else "Summary"
    summary.append(["Diagnostic job", result.job_id])
    summary.append(["Kind", result.kind])
    summary.append(["Completed at", result.completed_at])
    summary.append(["Model rank", result.model.rank])
    for key, value in result.sample.items():
        summary.append([key, value])
    style_sheet(summary)
    config = workbook.create_sheet(
        "配置" if result.spec.language.value == "zh" else "Configuration"
    )
    for key, value in result.spec.model_dump(mode="json").items():
        config.append([key, stringify(value)])
    style_sheet(config)
    for event in result.event_studies:
        sheet = workbook.create_sheet(f"Event-{event.method.value}"[:31])
        sheet.append(
            [
                "event_time",
                "estimate",
                "std_error",
                "conf_low",
                "conf_high",
                "weight",
                "identifiable",
            ]
        )
        for point in event.points:
            sheet.append(
                [
                    point.event_time,
                    point.estimate,
                    point.std_error,
                    point.conf_low,
                    point.conf_high,
                    point.weight,
                    point.identifiable,
                ]
            )
        sheet.append([])
        sheet.append(["pretrend_statistic", event.pretrend_statistic])
        sheet.append(["pretrend_df", event.pretrend_df])
        sheet.append(["pretrend_p_value", event.pretrend_p_value])
        style_sheet(sheet, header=True)
    for placebo in result.placebos:
        sheet = workbook.create_sheet(f"Placebo-{placebo.method.value}"[:31])
        sheet.append(["actual_att", placebo.actual_att])
        sheet.append(["mean", placebo.mean])
        sheet.append(["std_dev", placebo.std_dev])
        sheet.append(["quantile_low", placebo.quantile_low])
        sheet.append(["quantile_high", placebo.quantile_high])
        sheet.append(["empirical_p_value", placebo.empirical_p_value])
        sheet.append(["attempted", placebo.attempted])
        sheet.append(["failed", placebo.failed])
        sheet.append([])
        sheet.append(["iteration", "estimate"])
        for draw in placebo.draws:
            sheet.append([draw.iteration, draw.estimate])
        style_sheet(sheet)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_diagnostic_word(result: DiagnosticResult) -> bytes:
    language = result.spec.language.value
    document = Document()
    title = document.add_heading(
        "检验与图表报告" if language == "zh" else "Diagnostics and figures", level=0
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "注：所有图均报告95%置信区间，基准期为政策前1期。"
        if language == "zh"
        else "Notes: All event-study figures report 95% confidence intervals and omit period -1."
    )
    for event in result.event_studies:
        document.add_heading(
            "平行趋势检验" if language == "zh" else "Event-study estimates", level=1
        )
        image = BytesIO(figure_bytes(result, f"event_{event.method.value}", "png"))
        document.add_picture(image, width=Inches(6.3))
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, text in zip(
            table.rows[0].cells, ["Period", "Estimate", "SE", "Lower", "Upper"], strict=True
        ):
            cell.text = text
        for point in event.points:
            cells = table.add_row().cells
            for cell, value in zip(
                cells,
                [
                    point.event_time,
                    point.estimate,
                    point.std_error,
                    point.conf_low,
                    point.conf_high,
                ],
                strict=True,
            ):
                cell.text = format_number(value)
    for placebo in result.placebos:
        document.add_heading(placebo_title(placebo.method.value, language), level=1)
        image = BytesIO(figure_bytes(result, f"placebo_{placebo.method.value}", "png"))
        document.add_picture(image, width=Inches(6.3))
        summary = (
            f"ATT={placebo.actual_att:.6g}; empirical p={placebo.empirical_p_value:.4f}; "
            f"valid repetitions={len(placebo.draws)}."
        )
        document.add_paragraph(summary)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def style_sheet(sheet: Any, header: bool = False) -> None:
    navy = "173B57"
    if header and sheet.max_row:
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
    else:
        sheet["A1"].font = Font(bold=True, color=navy)
    for column in sheet.columns:
        width = min(max((len(str(cell.value or "")) for cell in column), default=8) + 2, 60)
        sheet.column_dimensions[column[0].column_letter].width = width


def stringify(value: Any) -> str:
    if isinstance(value, list):
        return ", ".join(map(str, value))
    if isinstance(value, dict):
        return "; ".join(f"{key}={item}" for key, item in value.items())
    return "" if value is None else str(value)


def format_number(value: Any) -> str:
    if value is None:
        return "—"
    if isinstance(value, float):
        return f"{value:.8g}"
    return str(value)
