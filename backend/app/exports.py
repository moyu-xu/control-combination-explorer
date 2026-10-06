from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import AnalysisResult, Language, ModelResult

LABELS = {
    "zh": {
        "title": "控制变量组合筛选结果",
        "summary": "运行摘要",
        "config": "运行配置",
        "classification": "控制变量分类",
        "top": "Top N 组合",
        "best": "最佳模型",
        "failures": "排除原因",
        "sample": "样本处理",
        "variable": "变量",
        "estimate": "系数",
        "se": "标准误",
        "stat": "t 值",
        "p": "p 值",
        "low": "95% 下限",
        "high": "95% 上限",
        "controls": "候选控制组合",
        "observations": "观测数",
        "rank": "排名",
        "stata": "Stata 复现命令",
        "raw_note": "所有 p 值均为未经多重检验校正的双侧 p 值。",
        "no_result": "没有满足指定方向和 5% 显著性门槛的组合。",
        "level": "层级",
        "dimension": "经济维度",
        "confidence": "置信度",
        "reason": "匹配依据",
        "manual": "人工修改",
    },
    "en": {
        "title": "Control-Variable Combination Search",
        "summary": "Run summary",
        "config": "Configuration",
        "classification": "Control classifications",
        "top": "Top N combinations",
        "best": "Best model",
        "failures": "Exclusion reasons",
        "sample": "Sample processing",
        "variable": "Variable",
        "estimate": "Estimate",
        "se": "Std. error",
        "stat": "t statistic",
        "p": "p value",
        "low": "95% lower",
        "high": "95% upper",
        "controls": "Candidate-control combination",
        "observations": "Observations",
        "rank": "Rank",
        "stata": "Stata reproduction command",
        "raw_note": "All p values are unadjusted, two-sided p values.",
        "no_result": "No combination met the requested direction and 5% significance threshold.",
        "level": "Level",
        "dimension": "Economic dimension",
        "confidence": "Confidence",
        "reason": "Match rationale",
        "manual": "Manually edited",
    },
}


def labels(result: AnalysisResult) -> dict[str, str]:
    language = (
        result.spec.language.value
        if isinstance(result.spec.language, Language)
        else str(result.spec.language)
    )
    return LABELS[language]


def build_excel(result: AnalysisResult) -> bytes:
    text = labels(result)
    workbook = Workbook()
    summary = workbook.active
    summary.title = text["summary"]
    summary.append([text["title"]])
    summary.append(["job_id", result.job_id])
    summary.append(["completed_at", result.completed_at])
    summary.append(["total_combinations", result.total_combinations])
    summary.append(["successful_models", result.successful_models])
    summary.append(["excluded_models", result.excluded_models])
    summary.append([])
    summary.append([text["raw_note"]])
    style_sheet(summary)

    config = workbook.create_sheet(text["config"])
    for key, value in result.spec.model_dump(mode="json").items():
        config.append([key, stringify(value)])
    style_sheet(config)

    classification = workbook.create_sheet(text["classification"])
    classification.append(
        [
            text["variable"],
            text["level"],
            text["dimension"],
            text["confidence"],
            text["reason"],
            text["manual"],
        ]
    )
    for item in result.spec.control_classifications:
        classification.append(
            [
                item.variable,
                item.level.value,
                item.dimension,
                item.confidence.value,
                item.reason,
                item.manually_modified,
            ]
        )
    style_sheet(classification, header=True)

    sample = workbook.create_sheet(text["sample"])
    for key, value in result.sample.items():
        sample.append([key, stringify(value)])
    style_sheet(sample)

    top = workbook.create_sheet(text["top"])
    top.append(
        [
            text["rank"],
            text["controls"],
            text["estimate"],
            text["se"],
            text["stat"],
            text["p"],
            text["observations"],
        ]
    )
    for item in result.top_models:
        top.append(
            [
                item.rank,
                ", ".join(item.controls),
                item.core_estimate,
                item.core_std_error,
                item.core_statistic,
                item.core_p_value,
                item.observations,
            ]
        )
    style_sheet(top, header=True)

    best = workbook.create_sheet(text["best"])
    best.append(
        [
            text["variable"],
            text["estimate"],
            text["se"],
            text["stat"],
            text["p"],
            text["low"],
            text["high"],
        ]
    )
    if result.best_model:
        append_coefficients(best, result.best_model)
        best.append([])
        best.append([text["stata"], result.stata_command])
    else:
        best.append([text["no_result"]])
    style_sheet(best, header=True)

    failures = workbook.create_sheet(text["failures"])
    failures.append(["reason", "count"])
    for reason, count in sorted(result.failure_counts.items()):
        failures.append([reason, count])
    style_sheet(failures, header=True)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def append_coefficients(sheet: Any, model: ModelResult) -> None:
    for coefficient in model.coefficients:
        sheet.append(
            [
                coefficient.variable,
                coefficient.estimate,
                coefficient.std_error,
                coefficient.statistic,
                coefficient.p_value,
                coefficient.conf_low,
                coefficient.conf_high,
            ]
        )


def style_sheet(sheet: Any, header: bool = False) -> None:
    navy = "173B57"
    if header and sheet.max_row:
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
    else:
        sheet["A1"].font = Font(bold=True, color=navy, size=14)
    for column in sheet.columns:
        width = min(max((len(str(cell.value or "")) for cell in column), default=8) + 2, 60)
        sheet.column_dimensions[column[0].column_letter].width = width
    sheet.freeze_panes = "A2" if header else None


def build_word(result: AnalysisResult) -> bytes:
    text = labels(result)
    document = Document()
    title = document.add_heading(text["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(text["raw_note"])

    document.add_heading(text["summary"], level=1)
    summary = document.add_table(rows=0, cols=2)
    summary.style = "Table Grid"
    for key, value in [
        ("Total combinations", result.total_combinations),
        ("Successful models", result.successful_models),
        ("Excluded models", result.excluded_models),
        ("Common-sample observations", result.sample.get("common_rows")),
        ("Raw combinations", result.sample.get("raw_combination_count")),
        ("Dimension-excluded combinations", result.sample.get("dimension_excluded_count")),
        ("Classification rule version", result.spec.classification_rule_version),
    ]:
        cells = summary.add_row().cells
        cells[0].text = key
        cells[1].text = stringify(value)

    document.add_heading(text["classification"], level=1)
    classification_table = document.add_table(rows=1, cols=5)
    classification_table.style = "Table Grid"
    for cell, heading in zip(
        classification_table.rows[0].cells,
        [text["variable"], text["level"], text["dimension"], text["confidence"], text["reason"]],
        strict=True,
    ):
        cell.text = heading
    for item in result.spec.control_classifications:
        cells = classification_table.add_row().cells
        for cell, value in zip(
            cells,
            [item.variable, item.level.value, item.dimension, item.confidence.value, item.reason],
            strict=True,
        ):
            cell.text = stringify(value)

    if not result.best_model:
        document.add_paragraph(text["no_result"])
    else:
        document.add_heading(text["best"], level=1)
        document.add_paragraph(f"{text['controls']}: {', '.join(result.best_model.controls)}")
        table = document.add_table(rows=1, cols=7)
        table.style = "Table Grid"
        headings = [
            text["variable"],
            text["estimate"],
            text["se"],
            text["stat"],
            text["p"],
            text["low"],
            text["high"],
        ]
        for cell, heading in zip(table.rows[0].cells, headings, strict=True):
            cell.text = heading
        for coefficient in result.best_model.coefficients:
            cells = table.add_row().cells
            values = [
                coefficient.variable,
                coefficient.estimate,
                coefficient.std_error,
                coefficient.statistic,
                coefficient.p_value,
                coefficient.conf_low,
                coefficient.conf_high,
            ]
            for cell, value in zip(cells, values, strict=True):
                cell.text = format_number(value)
        document.add_heading(text["stata"], level=2)
        document.add_paragraph(result.stata_command or "")

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def stringify(value: Any) -> str:
    if isinstance(value, list):
        return ", ".join(map(str, value))
    if isinstance(value, dict):
        return "; ".join(f"{key}={item}" for key, item in value.items())
    return "" if value is None else str(value)


def format_number(value: Any) -> str:
    if isinstance(value, float):
        return f"{value:.8g}"
    return stringify(value)
